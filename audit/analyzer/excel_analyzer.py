import openpyxl

class ExcelAnalyzer:
    def __init__(self):
        pass

    def analyze(self, filepath):
        try:
            # We need data_only=False to see formulas
            wb = openpyxl.load_workbook(filepath, data_only=False, read_only=True)
            sheets = {}
            sheet_order = wb.sheetnames
            
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                
                # Column headers and order
                headers = []
                # ws[1] gives a tuple of cells in the first row
                # In read_only mode, we should be careful with ws[1] if it's empty
                try:
                    for cell in ws[1]:
                        val = str(cell.value) if cell.value is not None else ""
                        headers.append(val)
                except Exception:
                    pass # Handle empty sheets
                
                # Track formulas
                formulas = {}
                # In read_only mode, iterating over cells can be slow for large files
                # But we need it for formulas. We'll limit the range if possible, 
                # but for static audit we usually care about the core data area.
                # Let's check max_row/max_column and maybe cap it for safety
                max_r = min(ws.max_row or 0, 1000) # Cap for performance in audit
                max_c = min(ws.max_column or 0, 50)
                
                for r in range(1, max_r + 1):
                    for c in range(1, max_c + 1):
                        cell = ws.cell(row=r, column=c)
                        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
                            formulas[f"{openpyxl.utils.get_column_letter(c)}{r}"] = cell.value
                
                sheets[sheet_name] = {
                    "columns": headers,
                    "column_count": len(headers),
                    "row_count": ws.max_row,
                    "formulas": formulas,
                    "formula_count": len(formulas),
                    "merged_cells": [str(r) for r in ws.merged_cells.ranges]
                }
            
            return {
                "sheets": sheets,
                "sheet_order": sheet_order,
                "type": "excel"
            }
        except Exception as e:
            return {"error": str(e), "type": "excel"}
